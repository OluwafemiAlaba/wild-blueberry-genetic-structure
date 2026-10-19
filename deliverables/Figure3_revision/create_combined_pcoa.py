from pathlib import Path
import openpyxl,numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
R=Path(__file__).resolve().parents[2];O=Path(__file__).resolve().parent
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'pdf.fonttype':42,'svg.fonttype':'none'})
colors={'Deblois':'#6A51A3','Jonesboro':'#4D4D4D','Hope':'#3A923A','Fort Kent':'#E69F00'}
rename={'Wyman':'Deblois','BBHF':'Jonesboro','FortKent':'Fort Kent'}
fig,axs=plt.subplots(1,2,figsize=(10,4.8));fig.subplots_adjust(left=.085,right=.98,top=.93,bottom=.24,wspace=.30)
for ax,name,letter in zip(axs,['angust_geographic','myrt_without_hope'],'AB'):
 w=openpyxl.load_workbook(R/'inputs/amova'/f'{name}.xlsx',read_only=True,data_only=True);s=w['PCoA'];n=int(s['B5'].value)
 rows=list(s.iter_rows(min_row=33,max_row=32+n,max_col=3,values_only=True));src=list(w['Sheet1'].iter_rows(min_row=4,max_row=n+3,max_col=2,values_only=True))
 def norm(x):
  for a,b in rename.items():x=x.replace(a,b)
  return x
 assert [norm(r[0]) for r in rows]==[norm(r[1]) for r in src]
 for group in dict.fromkeys(norm(r[0]) for r in rows):
  sub=[r for r in rows if norm(r[0])==group];ax.scatter([r[1] for r in sub],[r[2] for r in sub],s=34,c=colors[group],edgecolor='white',lw=.4,marker={'Deblois':'o','Hope':'s','Jonesboro':'^','Fort Kent':'D'}[group],label=group)
 ax.axhline(0,lw=.6,c='#dddddd',zorder=0);ax.axvline(0,lw=.6,c='#dddddd',zorder=0)
 ax.set_xlabel(f'Coordinate 1 ({s["B10"].value:.2f}%)');ax.set_ylabel(f'Coordinate 2 ({s["C10"].value:.2f}%)');ax.set_aspect('equal',adjustable='datalim');ax.margins(.12)
 ax.text(-.14,1.04,letter,transform=ax.transAxes,fontsize=18,weight='bold');ax.legend(loc='upper center',bbox_to_anchor=(.5,-.18),ncol=3,frameon=False,fontsize=9)
 for sp in ax.spines.values():
  sp.set_visible(True);sp.set_linewidth(.8);sp.set_color('#333333')
 print(name,n,s['B10'].value,s['C10'].value)
for ext in ['png','pdf','svg']:fig.savefig(O/f'Figure3_combined_PCoA.{ext}',dpi=400,facecolor='white')
fig.savefig(O/'Figure3_preview.png',dpi=150,facecolor='white')
